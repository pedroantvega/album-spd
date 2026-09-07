# -*- coding: utf-8 -*-
"""
docx_builder.py — Genera el DOCX del álbum (versión editable, aproximada al PDF).
"""
import os
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

import imagenes as imgmod
from datos import resumen_una_linea
from diseno import TITULO_ALBUM, BOX_P_W, BOX_P_H, BOX_MG_LONG, BOX_MG_SHORT
from reportlab.lib.units import mm as MM


def _mm(v_pt):
    """v_pt viene en unidades reportlab (puntos); convertir a mm reales."""
    return v_pt / MM


def set_page_border(section, color='000000', size=24):
    """Añade un marco grueso a toda la página (borde de sección)."""
    sectPr = section._sectPr
    pgBorders = OxmlElement('w:pgBorders')
    pgBorders.set(qn('w:offsetFrom'), 'page')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(size))
        el.set(qn('w:space'), '18')
        el.set(qn('w:color'), color)
        pgBorders.append(el)
    sectPr.append(pgBorders)


def add_field(paragraph, field_code):
    run = paragraph.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = field_code
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def setup_header_footer(doc, anio):
    section = doc.sections[0]
    section.page_height = Mm(297)
    section.page_width = Mm(210)
    section.top_margin = Mm(22)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    set_page_border(section)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = hp.add_run(TITULO_ALBUM)
    run.bold = True
    run.font.size = Pt(13)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run(f'{anio}  —  Página ')
    r1.font.size = Pt(9)
    add_field(fp, 'PAGE')
    r2 = fp.add_run(' de ')
    r2.font.size = Pt(9)
    add_field(fp, 'NUMPAGES')


def add_caption(doc, n, entry):
    p = doc.add_paragraph()
    r = p.add_run(f"Nº {n}   Edifil {entry['edifil']}")
    r.bold = True
    r.font.size = Pt(10.5)
    p2 = doc.add_paragraph()
    r2 = p2.add_run(f"{entry['serie']}  —  {entry['fecha']}")
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    texto = entry.get('resumen_texto') or ''
    if texto:
        p3 = doc.add_paragraph()
        r3 = p3.add_run(texto)
        r3.italic = True
        r3.font.size = Pt(8.3)
        r3.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def add_imagen_con_marco(doc, ruta, ancho_mm, alto_mm, cache_dir, rotar=False, escala_pct=100):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '8')
        el.set(qn('w:color'), '999999')
        borders.append(el)
    tblPr.append(borders)

    cell.width = Mm(ancho_mm)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if ruta and os.path.exists(ruta):
        img_final = imgmod.preparar_imagen(ruta, cache_dir, rotar_90=rotar)
        with Image.open(img_final) as im:
            iw, ih = im.size
        ratio = iw / ih
        factor = escala_pct / 100.0
        if ancho_mm / alto_mm > ratio:
            h_mm = alto_mm * 0.97 * factor
            w_mm = h_mm * ratio
        else:
            w_mm = ancho_mm * 0.97 * factor
            h_mm = w_mm / ratio
        # Word no recorta automáticamente si la imagen supera la celda;
        # limitamos el factor para que no se salga de forma descontrolada.
        w_mm = min(w_mm, ancho_mm * 1.15)
        h_mm = min(h_mm, alto_mm * 1.15)
        run = p.add_run()
        run.add_picture(img_final, width=Mm(w_mm), height=Mm(h_mm))
    else:
        run = p.add_run('(sin imagen disponible)')
        run.italic = True

    for row in table.rows:
        row.height = Mm(alto_mm)
    return table


def construir_docx(anio, grupos, cache_dir, ruta_salida, lineas_resumen=1, escala_imagen=100):
    pequenos = [g for g in grupos if g['tamano'] == 'P']
    grandes = [g for g in grupos if g['tamano'] == 'MG']
    extra = [g for g in grupos if g['tamano'] == 'EXTRA']
    revisar = [g for g in grupos if g['tamano'] not in ('P', 'MG', 'EXTRA')]
    pequenos.extend(revisar)

    from datos import resumen_texto
    max_chars = 100 if lineas_resumen <= 1 else 220
    for g in grupos:
        g['resumen_texto'] = ''
    for g in pequenos + grandes:
        if g['descripciones']:
            g['resumen_texto'] = resumen_texto(g['descripciones'][0], max_chars=max_chars)

    numero_map = {id(g): i + 1 for i, g in enumerate(grupos)}

    doc = Document()
    setup_header_footer(doc, anio)

    # --- Índice ---
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rh = h.add_run(str(anio))
    rh.bold = True
    rh.font.size = Pt(28)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run('Sobres de Primer Día — España — Índice del año')
    rs.font.size = Pt(11)
    rs.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    tabla = doc.add_table(rows=1, cols=8)
    tabla.style = 'Light Grid Accent 1'
    hdr = tabla.rows[0].cells
    for i, txt in enumerate(['Nº', 'Edifil', 'Fecha', 'Serie / Emisión', 'Cant.', 'Valor Edifil', 'Total', 'Tengo']):
        hdr[i].text = txt
    total_sellos = 0
    for i, g in enumerate(grupos):
        cantidad = len(str(g['edifil']).split(','))
        total_sellos += cantidad
        row = tabla.add_row().cells
        row[0].text = str(i + 1)
        row[1].text = g['edifil']
        row[2].text = g['fecha']
        row[3].text = g['serie']
        row[4].text = str(cantidad)
        row[5].text = ''
        row[6].text = ''
        row[7].text = ''
    fila_total = tabla.add_row().cells
    fila_total[3].text = 'Total del año'
    fila_total[3].paragraphs[0].runs[0].bold = True
    fila_total[4].text = str(total_sellos)
    fila_total[4].paragraphs[0].runs[0].bold = True

    nota = doc.add_paragraph()
    rn = nota.add_run('"Valor Edifil" y "Total" quedan en blanco para rellenar a mano '
                       '(o desde tu propio catálogo cuando lo tengas centralizado).')
    rn.italic = True
    rn.font.size = Pt(8.5)
    rn.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_page_break()

    # --- Sobres pequeños, 2 por página ---
    for i in range(0, len(pequenos), 2):
        for g in pequenos[i:i + 2]:
            add_imagen_con_marco(doc, g['imagen_ruta'], _mm(BOX_P_W), _mm(BOX_P_H), cache_dir,
                                  escala_pct=escala_imagen)
            add_caption(doc, numero_map[id(g)], g)
            doc.add_paragraph()
        doc.add_page_break()

    # --- Sobres medianos/grandes, 1 por página (tamaño real cuando se conoce) ---
    for g in grandes:
        if g.get('tamano_real_mm'):
            ancho_real, alto_real = g['tamano_real_mm']
        else:
            ancho_real, alto_real = _mm(BOX_MG_SHORT), _mm(BOX_MG_LONG)
        add_imagen_con_marco(doc, g['imagen_ruta'], ancho_real, alto_real, cache_dir, rotar=True,
                              escala_pct=escala_imagen)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"Nº {numero_map[id(g)]}   Edifil {g['edifil']}  —  {ancho_real:.0f}×{alto_real:.0f}mm")
        r.bold = True
        r.font.size = Pt(10.5)
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(f"{g['serie']}  —  {g['fecha']}")
        r2.font.size = Pt(9.5)
        doc.add_page_break()

    # --- Sobres extra grande: solo datos, sin imagen ---
    for g in extra:
        ancho_real, alto_real = g.get('tamano_real_mm') or (None, None)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"Nº {numero_map[id(g)]}   Edifil {g['edifil']}")
        r.bold = True
        r.font.size = Pt(16)
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(g['serie'])
        r2.font.size = Pt(12)
        p3 = doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.add_run(g['fecha']).font.size = Pt(11)
        if ancho_real:
            p4 = doc.add_paragraph()
            p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r4 = p4.add_run(f'Sobre extra grande — {ancho_real:.0f}×{alto_real:.0f}mm')
            r4.italic = True
            r4.font.size = Pt(9.5)
        p5 = doc.add_paragraph()
        p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r5 = p5.add_run('(sin foto en esta página — sobre demasiado grande para la plantilla '
                         'estándar; colócalo aparte con estos datos como referencia)')
        r5.italic = True
        r5.font.size = Pt(9)
        r5.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
        doc.add_page_break()

    # --- Descripciones completas ---
    dt = doc.add_paragraph()
    rt = dt.add_run('Descripciones completas')
    rt.bold = True
    rt.font.size = Pt(15)
    for g in grupos:
        for desc in g['descripciones']:
            pt = doc.add_paragraph()
            rpt = pt.add_run(f"Edifil {g['edifil']} — {g['serie']}")
            rpt.bold = True
            rpt.font.size = Pt(10)
            pb = doc.add_paragraph()
            rpb = pb.add_run(desc)
            rpb.font.size = Pt(9)
            rpb.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    doc.save(ruta_salida)
